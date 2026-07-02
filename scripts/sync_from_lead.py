#!/usr/bin/env python3
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import http.cookiejar
import importlib
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


PROJECT_DIR = Path(__file__).resolve().parents[1]
STATE_PATH = PROJECT_DIR / "sync_state.json"
ENV_PATH = PROJECT_DIR / ".env"
AUTH_DIR = PROJECT_DIR / ".auth"
AUTH_STATE_PATH = AUTH_DIR / "lead_state.json"
# The sync-floor date is not hardcoded — it is read from the DB `min_sync_date`
# setting (editable in the admin page). Empty/unset means "no floor".


def load_db_store():
    try:
        if str(PROJECT_DIR) not in sys.path:
            sys.path.insert(0, str(PROJECT_DIR))
        return importlib.import_module("scripts.db_store"), None
    except Exception as exc:  # pragma: no cover
        return None, exc


def load_env(path: Path) -> dict[str, str]:
    env: dict[str, str] = {k: v for k, v in os.environ.items() if k.startswith("LEAD_")}
    if not path.exists():
        return env
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def norm(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip())


def money(v: Any) -> float:
    m = re.search(r"-?\d+(?:\.\d+)?", norm(v).replace(",", ""))
    return float(m.group()) if m else 0.0


def parse_date_text(value: Any) -> dt.date | None:
    text = norm(value)
    if not text:
        return None
    match = re.search(r"(20\d{2})[/-](\d{2})[/-](\d{2})", text)
    if match:
        return dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    match = re.search(r"(\d{2})/(\d{2})", text)
    if match:
        month = int(match.group(1))
        day = int(match.group(2))
        year = 2026
        return dt.date(year, month, day)
    return None


def row_date_at_or_after(row: list[Any], indexes: tuple[int, ...], minimum: dt.date | None) -> bool:
    # No configured floor → keep every row (the DB upsert is idempotent).
    if minimum is None:
        return True
    for idx in indexes:
        if idx < len(row):
            parsed = parse_date_text(row[idx])
            if parsed and parsed >= minimum:
                return True
    return False


def load_auth_state() -> dict[str, Any]:
    if AUTH_STATE_PATH.exists():
        try:
            return json.loads(AUTH_STATE_PATH.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_auth_state(state: dict[str, Any]) -> None:
    AUTH_DIR.mkdir(parents=True, exist_ok=True)
    AUTH_STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def load_state() -> dict[str, Any]:
    if STATE_PATH.exists():
        try:
            return json.loads(STATE_PATH.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_state(state: dict[str, Any]) -> None:
    STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


class TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_table = False
        self.in_row = False
        self.in_cell = False
        self.tables: list[list[list[str]]] = []
        self.current_table: list[list[str]] = []
        self.current_row: list[str] = []
        self.buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.in_table = True
            self.current_table = []
        elif self.in_table and tag == "tr":
            self.in_row = True
            self.current_row = []
        elif self.in_row and tag in ("td", "th"):
            self.in_cell = True
            self.buf = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.in_cell:
            self.current_row.append(norm("".join(self.buf)))
            self.in_cell = False
        elif tag == "tr" and self.in_row:
            if any(self.current_row):
                self.current_table.append(self.current_row[:])
            self.in_row = False
        elif tag == "table" and self.in_table:
            if self.current_table:
                self.tables.append(self.current_table[:])
            self.in_table = False

    def handle_data(self, data):
        if self.in_cell:
            self.buf.append(data)


def parse_tables(html: str) -> list[list[list[str]]]:
    parser = TableParser()
    parser.feed(html)
    return parser.tables


def data_key(row: list[Any], idxs: tuple[int, ...]) -> str:
    return " | ".join(norm(row[i]) if i < len(row) else "" for i in idxs)


class FormParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.forms: list[dict[str, Any]] = []
        self._in_form = False
        self._form: dict[str, Any] = {}

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "form":
            self._in_form = True
            self._form = {"method": attrs.get("method", "get").upper(), "action": attrs.get("action", ""), "inputs": [], "buttons": []}
        elif self._in_form and tag == "input":
            self._form["inputs"].append(attrs)
        elif self._in_form and tag == "button":
            self._form["buttons"].append(attrs)

    def handle_endtag(self, tag):
        if tag == "form" and self._in_form:
            self.forms.append(self._form)
            self._in_form = False


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            attrs = dict(attrs)
            href = attrs.get("href")
            if href:
                self.links.append(href)


HTTP_RETRIES = 3
HTTP_BACKOFF_SECONDS = 1.5


def http_get_meta(opener, url: str) -> tuple[str, int, str]:
    # Retry transient failures with backoff so one flaky request doesn't silently
    # drop a date window's rows during the complete scrape.
    last_exc: Exception | None = None
    for attempt in range(HTTP_RETRIES):
        try:
            req = urllib.request.Request(url)
            with opener.open(req, timeout=60) as resp:
                html = resp.read().decode("utf-8", errors="ignore")
                final_url = getattr(resp, "geturl", lambda: url)()
                status = int(getattr(resp, "status", resp.getcode()))
                return final_url, status, html
        except Exception as exc:  # noqa: BLE001 — retry any transient network/HTTP error
            last_exc = exc
            if attempt < HTTP_RETRIES - 1:
                time.sleep(HTTP_BACKOFF_SECONDS * (attempt + 1))
    raise last_exc if last_exc is not None else RuntimeError(f"GET failed: {url}")


def http_get(opener, url: str) -> str:
    _, _, html = http_get_meta(opener, url)
    return html


def safe_http_get(opener, url: str) -> str:
    """GET that never raises — returns "" after exhausting retries. Used for
    optional pages so one page hiccup can't abort the whole sync."""
    try:
        return http_get(opener, url)
    except Exception:
        return ""


def http_get_bytes(opener, url: str) -> bytes:
    """Binary GET (for invoice .xlsx downloads), with the same retry/backoff."""
    last_exc: Exception | None = None
    for attempt in range(HTTP_RETRIES):
        try:
            with opener.open(urllib.request.Request(url), timeout=60) as resp:
                return resp.read()
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            if attempt < HTTP_RETRIES - 1:
                time.sleep(HTTP_BACKOFF_SECONDS * (attempt + 1))
    raise last_exc if last_exc is not None else RuntimeError(f"GET failed: {url}")


# Invoice .xlsx columns (admin/invoices.php downloads) -> our field names. These
# carry the EXACT per-shipment economics: Policy Price (= the stored shipping
# charge), Base Cost (actual carrier cost), Over Fee, COD Fixed.
_INVOICE_COLS = {
    "id": "order_id", "policy price": "policy", "base cost": "base_cost",
    "over fee": "over_fee", "cod fixed": "cod_fixed", "status": "status",
}


def parse_invoice_workbook(data: bytes) -> list[dict[str, Any]]:
    """Parse one invoice .xlsx into per-shipment rows via openpyxl."""
    from io import BytesIO
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip().lower() if h is not None else "" for h in next(it)]
        except StopIteration:
            return []
        col = {_INVOICE_COLS[h]: i for i, h in enumerate(header) if h in _INVOICE_COLS}
        if "order_id" not in col:
            return []
        out: list[dict[str, Any]] = []
        for r in it:
            if not r:
                continue
            oid = r[col["order_id"]]
            if oid in (None, ""):
                continue
            out.append({
                "order_id": norm(oid).lstrip("#"),
                "base_cost": money(r[col["base_cost"]]) if "base_cost" in col else 0.0,
                "policy": money(r[col["policy"]]) if "policy" in col else 0.0,
                "over_fee": money(r[col["over_fee"]]) if "over_fee" in col else 0.0,
                "cod_fixed": money(r[col["cod_fixed"]]) if "cod_fixed" in col else 0.0,
                "status": norm(r[col["status"]]) if "status" in col else "",
            })
        return out
    finally:
        wb.close()


def collect_invoice_costs(opener, base: str) -> dict[str, dict[str, Any]]:
    """Download every invoice .xlsx and map order_id -> per-shipment actuals.
    Invoices don't overlap, so each shipment appears at most once."""
    base = base.rstrip("/")
    html = safe_http_get(opener, f"{base}/admin/invoices.php?date_from=2024-01-01&date_to=2031-12-31")
    urls = sorted(set(re.findall(r"/uploads/invoices/[^\s\"'<>]+\.xlsx", html)))
    costs: dict[str, dict[str, Any]] = {}
    for u in urls:
        try:
            data = http_get_bytes(opener, f"{base}{u}")
            for row in parse_invoice_workbook(data):
                costs[row["order_id"]] = row
        except Exception as exc:  # noqa: BLE001
            print(f"[sync] invoice parse skipped {u}: {exc}", file=sys.stderr)
    return costs


def http_post_meta(opener, url: str, data: dict[str, str]) -> tuple[str, int, str]:
    body = urllib.parse.urlencode(data).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers={"Content-Type": "application/x-www-form-urlencoded"})
    with opener.open(req, timeout=60) as resp:
        html = resp.read().decode("utf-8", errors="ignore")
        final_url = getattr(resp, "geturl", lambda: url)()
        status = int(getattr(resp, "status", resp.getcode()))
        return final_url, status, html


def log_page_sample(logs: list[dict[str, Any]], label: str, requested_url: str, final_url: str, status: int, html: str, reason: str = "") -> None:
    snippet = html[:500]
    logs.append({
        "kind": "page_sample",
        "page": label,
        "requested_url": requested_url,
        "final_url": final_url,
        "status": status,
        "reason": reason,
        "html_500": snippet,
    })


def scrape_site(env: dict[str, str]) -> dict[str, Any]:
    base = env["LEAD_BASE_URL"].rstrip("/")
    username = env["LEAD_USERNAME"]
    password = env["LEAD_PASSWORD"]
    cj = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    logs: list[dict[str, Any]] = []
    force_login = os.environ.get("LEAD_FORCE_LOGIN_EACH_RUN", "1").strip().lower() not in ("0", "false", "no")
    auth_state = {} if force_login else load_auth_state()
    if force_login:
        try:
            AUTH_STATE_PATH.unlink(missing_ok=True)
        except Exception:
            pass

    def page_looks_logged_in(html: str) -> tuple[bool, dict[str, bool]]:
        parser = LinkParser()
        parser.feed(html)
        lower = html.lower()
        checks = {
            "has_admin_links": any(token in html for token in ("/admin/shipments.php", "/admin/wallet.php", "/admin/collect-cod.php", "/admin/dashboard.php")),
            "shipments_link": any("/admin/shipments.php" in link for link in parser.links),
            "wallet_link": any("/admin/wallet.php" in link for link in parser.links),
            "cod_link": any("/admin/collect-cod.php" in link for link in parser.links),
            "not_login_form": "password" not in lower or "/login" not in lower,
        }
        return all(checks.values()), checks

    def page_is_login_like(html: str, final_url: str) -> bool:
        lower = html.lower()
        return "login" in final_url.lower() or ("password" in lower and "username" in lower)

    def fetch_page(label: str, url: str) -> tuple[str, int, str]:
        final_url, status, html = http_get_meta(opener, url)
        reason = ""
        if final_url != url:
            reason = f"redirected_to={final_url}"
        elif page_is_login_like(html, final_url):
            reason = "login_page_detected"
        log_page_sample(logs, label, url, final_url, status, html, reason)
        return final_url, status, html

    if auth_state.get("cookie"):
        opener.addheaders = [("Cookie", auth_state["cookie"])]
        _, dashboard_status, cached_dashboard = fetch_page("dashboard.php", f"{base}/admin/dashboard.php")
        ok, cached_checks = page_looks_logged_in(cached_dashboard)
        logs.append({"kind": "probe", "page": "dashboard.php", "cached_session": ok, "checks": cached_checks})
        shipments_final_url, shipments_status, shipments_probe = fetch_page("shipments.php", f"{base}/admin/shipments.php")
        shipment_ok, shipment_checks = page_looks_logged_in(shipments_probe)
        logs.append({"kind": "probe", "page": "shipments.php", "cached_session": shipment_ok, "checks": shipment_checks, "final_url": shipments_final_url, "status": shipments_status})
        _, wallet_status, wallet_probe = fetch_page("wallet.php", f"{base}/admin/wallet.php")
        wallet_ok, wallet_checks = page_looks_logged_in(wallet_probe)
        logs.append({"kind": "probe", "page": "wallet.php", "cached_session": wallet_ok, "checks": wallet_checks, "status": wallet_status})
        _, cod_status, cod_probe = fetch_page("collect-cod.php", f"{base}/admin/collect-cod.php")
        cod_ok, cod_checks = page_looks_logged_in(cod_probe)
        logs.append({"kind": "probe", "page": "collect-cod.php", "cached_session": cod_ok, "checks": cod_checks, "status": cod_status})
        if ok and shipment_ok and wallet_ok and cod_ok:
            html = {
                "dashboard.php": cached_dashboard,
                "shipments.php": shipments_probe,
                "wallet.php": wallet_probe,
                "collect-cod.php": cod_probe,
                "sync-status.php": http_get(opener, f"{base}/admin/sync-status.php"),
                "shipping-companies.php": safe_http_get(opener, f"{base}/admin/shipping-companies.php"),
                "reports.php": safe_http_get(opener, f"{base}/admin/reports.php"),
                "pending-recharges.php": safe_http_get(opener, f"{base}/admin/pending-recharges.php"),
            }
            logs.append({"kind": "session", "source": "saved_state", "used": True})
            return {
                "opener": opener,
                "logs": logs,
                "html": html,
                "checks": {
                    "current_url": f"{base}/admin/dashboard.php",
                    "has_dashboard_links": ok,
                    "shipments_opened": shipment_ok,
                    "wallet_opened": wallet_ok,
                    "cod_opened": cod_ok,
                    "session_used": True,
                    "dashboard_checks": cached_checks,
                    "shipments_checks": shipment_checks,
                    "wallet_checks": wallet_checks,
                    "cod_checks": cod_checks,
                },
            }

    login_final_url, login_status, login_html = fetch_page("login.php", f"{base}/login.php")
    logs.append({"kind": "request", "url": f"{base}/login.php", "final_url": login_final_url, "status": login_status})
    parser = FormParser()
    parser.feed(login_html)
    form = parser.forms[0] if parser.forms else {"method": "POST", "action": "/login.php", "inputs": []}
    post_data: dict[str, str] = {}
    for inp in form["inputs"]:
        name = inp.get("name")
        if not name:
            continue
        typ = (inp.get("type") or "").lower()
        if typ == "password":
            post_data[name] = password
        elif typ in ("text", "email") or "user" in name.lower() or "login" in name.lower():
            post_data[name] = username
        elif typ == "hidden":
            post_data[name] = inp.get("value", "")
    for btn in form.get("buttons", []):
        btn_name = btn.get("name")
        if btn_name and btn.get("type", "").lower() in ("submit", ""):
            post_data.setdefault(btn_name, btn.get("value", "") or "1")
    action = form.get("action") or "/login.php"
    if not action.startswith("http"):
        action = f"{base}/{action.lstrip('/')}"
    login_post_final_url, login_post_status, login_result = http_post_meta(opener, action, post_data)
    login_result_html = login_result[:500]
    logs.append({
        "kind": "request",
        "url": action,
        "final_url": login_post_final_url,
        "status": login_post_status,
        "response_500": login_result_html,
    })

    dashboard_final_url, dashboard_status, dashboard_html = fetch_page("dashboard.php", f"{base}/admin/dashboard.php")
    logs.append({"kind": "request", "url": f"{base}/admin/dashboard.php", "final_url": dashboard_final_url, "status": dashboard_status})
    has_dashboard_links, dashboard_checks = page_looks_logged_in(dashboard_html)
    shipments_final_url, shipments_status, shipments_probe = fetch_page("shipments.php", f"{base}/admin/shipments.php")
    logs.append({"kind": "request", "url": f"{base}/admin/shipments.php", "final_url": shipments_final_url, "status": shipments_status})
    shipments_checks = {
        "status_200": shipments_status == 200,
        "redirected": shipments_final_url != f"{base}/admin/shipments.php",
        "login_like": page_is_login_like(shipments_probe, shipments_final_url),
    }
    shipments_ok = shipments_status == 200 and not shipments_checks["login_like"]
    wallet_final_url, wallet_status, wallet_probe = fetch_page("wallet.php", f"{base}/admin/wallet.php")
    logs.append({"kind": "request", "url": f"{base}/admin/wallet.php", "final_url": wallet_final_url, "status": wallet_status})
    wallet_checks = {
        "status_200": wallet_status == 200,
        "redirected": wallet_final_url != f"{base}/admin/wallet.php",
        "login_like": page_is_login_like(wallet_probe, wallet_final_url),
    }
    wallet_ok = wallet_status == 200 and not wallet_checks["login_like"]
    cod_final_url, cod_status, cod_probe = fetch_page("collect-cod.php", f"{base}/admin/collect-cod.php")
    logs.append({"kind": "request", "url": f"{base}/admin/collect-cod.php", "final_url": cod_final_url, "status": cod_status})
    cod_checks = {
        "status_200": cod_status == 200,
        "redirected": cod_final_url != f"{base}/admin/collect-cod.php",
        "login_like": page_is_login_like(cod_probe, cod_final_url),
    }
    cod_ok = cod_status == 200 and not cod_checks["login_like"]
    login_issue = None
    if not shipments_ok:
        login_issue = "shipments page was not accessible"
    elif not wallet_ok:
        login_issue = "wallet page was not accessible"
    elif not cod_ok:
        login_issue = "cod page was not accessible"
    if login_issue:
        logs.append({
            "kind": "login_issue",
            "page": "shipments.php" if not shipments_ok else "wallet.php" if not wallet_ok else "collect-cod.php",
            "reason": login_issue,
            "shipments": {"final_url": shipments_final_url, "status": shipments_status, "checks": shipments_checks},
            "wallet": {"final_url": wallet_final_url, "status": wallet_status, "checks": wallet_checks},
            "cod": {"final_url": cod_final_url, "status": cod_status, "checks": cod_checks},
        })

    saved_cookie = None
    for cookie in cj:
        if cookie.name == "PHPSESSID":
            saved_cookie = f"{cookie.name}={cookie.value}"
            break
    if saved_cookie:
        save_auth_state({"cookie": saved_cookie, "saved_at": dt.datetime.now().isoformat(timespec="seconds")})

    html: dict[str, str] = {
        "dashboard.php": dashboard_html,
        "shipments.php": shipments_probe,
        "wallet.php": wallet_probe,
        "collect-cod.php": cod_probe,
        "sync-status.php": http_get(opener, f"{base}/admin/sync-status.php"),
        "shipping-companies.php": safe_http_get(opener, f"{base}/admin/shipping-companies.php"),
                "reports.php": safe_http_get(opener, f"{base}/admin/reports.php"),
        "pending-recharges.php": safe_http_get(opener, f"{base}/admin/pending-recharges.php"),
    }
    logs.append({"kind": "request", "url": f"{base}/admin/sync-status.php"})
    return {
        "opener": opener,
        "logs": logs,
        "html": html,
        "checks": {
            "current_url": f"{base}/admin/dashboard.php",
            "has_dashboard_links": has_dashboard_links,
            "shipments_opened": shipments_ok,
            "wallet_opened": wallet_ok,
            "cod_opened": cod_ok,
            "session_used": False,
            "dashboard_checks": dashboard_checks,
            "shipments_checks": shipments_checks,
            "wallet_checks": wallet_checks,
            "cod_checks": cod_checks,
            "dashboard_status": dashboard_status,
            "shipments_status": shipments_status,
            "wallet_status": wallet_status,
            "cod_status": cod_status,
            "dashboard_final_url": dashboard_final_url,
            "shipments_final_url": shipments_final_url,
            "wallet_final_url": wallet_final_url,
            "cod_final_url": cod_final_url,
            "login_status": login_status,
            "login_final_url": login_final_url,
            "login_http_status": login_status,
            "login_redirected_to_login": "login" in login_final_url.lower(),
            "login_html_500": login_html[:500],
        },
    }


def normalize_headers(row: list[str]) -> list[str]:
    return [norm(v) for v in row]


def _net_of_vat(amount: float, vat_rate: float) -> float:
    return round(amount / (1 + vat_rate), 2) if vat_rate > -1 else round(amount, 2)


def _wallet_numeric_key(row: list[Any]) -> int | None:
    match = re.search(r"\d+", norm(row[0]) if row else "")
    return int(match.group()) if match else None


def _wallet_vat_base(desc: str) -> float | None:
    match = re.search(r"\(([^)]*)\)", desc)
    if not match:
        return None
    value = money(match.group(1))
    return value if value else None


def _clear_wallet_vat_customers(wallet_rows: list[list[Any]]) -> set[str]:
    deposits_by_customer: dict[str, list[tuple[int, float]]] = {}
    for row in wallet_rows:
        customer = norm(row[2]) if len(row) > 2 else ""
        tx_type = norm(row[3]) if len(row) > 3 else ""
        key = _wallet_numeric_key(row)
        amount = money(row[4]) if len(row) > 4 else 0.0
        if customer and key is not None and tx_type == "إيداع" and amount > 0:
            deposits_by_customer.setdefault(customer, []).append((key, amount))

    customers: set[str] = set()
    for row in wallet_rows:
        customer = norm(row[2]) if len(row) > 2 else ""
        tx_type = norm(row[3]) if len(row) > 3 else ""
        desc = norm(row[5]) if len(row) > 5 else ""
        key = _wallet_numeric_key(row)
        amount = money(row[4]) if len(row) > 4 else 0.0
        base = _wallet_vat_base(desc)
        if not customer or key is None or base is None:
            continue
        is_clear_vat_deduction = (
            tx_type == "admin_deduction"
            and "خصم ضريبة القيمة المضافة عن شحن رصيد" in desc
            and "مرفوض" not in desc
            and "اضافة رصيد" not in desc
            and "تكرار" not in desc
        )
        if not is_clear_vat_deduction:
            continue
        expected_deposit = round(base + amount, 2)
        linked_deposit = any(
            abs(deposit_amount - expected_deposit) <= 0.05 and (key - 2) <= deposit_key <= (key + 1)
            for deposit_key, deposit_amount in deposits_by_customer.get(customer, [])
        )
        if linked_deposit:
            customers.add(customer)
    return customers


def wallet_tax_agreement_customers(wallet_rows: list[list[Any]]) -> set[str]:
    return _clear_wallet_vat_customers(wallet_rows)


BANK_VAT_AGREEMENT_CUSTOMERS = {
    "احمد محمد بن عبدالله الرزق متجر مقام الصبا | للات الموسيقية",
}


def bank_vat_agreement_customers() -> set[str]:
    return {customer for customer in BANK_VAT_AGREEMENT_CUSTOMERS if customer}


def stored_wallet_tax_agreement_customers(conn) -> set[str]:
    customers: set[str] = set()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT transaction_key, transaction_date, user_name, transaction_type, amount, description
                   FROM wallet_transactions
                   WHERE user_name <> 'user_name'
                     AND (
                        transaction_type = 'إيداع'
                        OR (
                            transaction_type = 'admin_deduction'
                            AND position('ضريبة القيمة المضافة' in coalesce(description, '')) > 0
                            AND position('شحن رصيد' in coalesce(description, '')) > 0
                        )
                     )
                   ORDER BY transaction_date NULLS LAST, transaction_key"""
            )
            rows = []
            for row in cur.fetchall():
                rows.append([
                    row["transaction_key"],
                    row["transaction_date"],
                    row["user_name"],
                    row["transaction_type"],
                    row["amount"],
                    row["description"],
                ])
            customers.update(_clear_wallet_vat_customers(rows))
    except Exception as exc:
        print(f"[sync] stored wallet VAT customers skipped: {exc}", file=sys.stderr)
    return {c for c in customers if c}


def shipment_record(row: list[list[Any]], snap=None, invoice_costs=None,
                    tax_agreement_customers: set[str] | None = None,
                    vat_rate: float = 0.15) -> dict[str, Any]:
    record = {
        "order_id": norm(row[0]) if len(row) > 0 else "",
        "tracking_number": norm(row[1]) if len(row) > 1 else "",
        "merchant_name": norm(row[2]) if len(row) > 2 else "",
        "store_name": norm(row[3]) if len(row) > 3 else "",
        "customer_name": norm(row[4]) if len(row) > 4 else "",
        "city": norm(row[5]) if len(row) > 5 else "",
        "carrier": norm(row[10]) if len(row) > 10 else "",
        "payment_type": norm(row[9]) if len(row) > 9 else "",
        "status": norm(row[12]) if len(row) > 12 else "",
        "shipment_date": parse_date_text(row[13]) if len(row) > 13 else None,
        "delivery_date": parse_date_text(row[14]) if len(row) > 14 else None,
        "weight": money(row[11]) if len(row) > 11 else 0.0,
        "cod_amount": money(row[7]) if len(row) > 7 else 0.0,
        "shipping_charge": money(row[6]) if len(row) > 6 else 0.0,
        "source_row": None,
        "source_hash": norm(row[0]) if row else "",
    }
    # Per-shipment ACTUALS — Lead's original real-economics formula, with VAT
    # applied only after the same old inputs are calculated:
    #   true profit = adjusted charge - adjusted(Base Cost + Over Fee + COD Fixed)
    # Billed shipments take Base Cost / Over Fee / COD Fixed from the invoice.
    # Current un-invoiced shipments use Lead's live platform price, overweight
    # starts after 10 kg at 2/kg, and COD platform fee is 3/order.
    from scripts import pricing as _pr
    realized_excluded = _pr.EXCLUDED_STATUSES + ("مرتجع",)
    charge = record["shipping_charge"]
    weight = record["weight"]
    realized = record["status"] not in realized_excluded
    is_cod = "COD" in record["payment_type"]
    tax_agreement_customers = tax_agreement_customers or set()
    customer_revenue = charge if record["merchant_name"] in tax_agreement_customers else _net_of_vat(charge, vat_rate)
    inv = (invoice_costs or {}).get(record["order_id"])
    if inv:
        base_cost_gross = inv["base_cost"]
        extra_cost_gross = round(inv["over_fee"] + inv["cod_fixed"], 2)
        cost_source = "invoice"
    else:
        carrier_key = _pr.CARRIER_ALIASES.get(record["carrier"], record["carrier"])
        carrier = (snap.carriers.get(carrier_key) if snap else None) or {}
        platform_gross = _pr.money(carrier.get("platform_gross"))
        platform_net = _pr.money(carrier.get("platform_net"))
        if platform_gross or platform_net:
            base_cost_gross = platform_gross or round(platform_net * (1 + vat_rate), 2)
            extra_cost_gross = round(max(weight - 10.0, 0.0) * 2.0 + (3.0 if is_cod else 0.0), 2)
            cost_source = "computed"
        else:
            base_cost_gross = extra_cost_gross = 0.0
            cost_source = "unknown"
    base_cost = _net_of_vat(base_cost_gross, vat_rate)
    extra_cost = _net_of_vat(extra_cost_gross, vat_rate)
    counted = realized and cost_source != "unknown"
    total_cost = round(base_cost + extra_cost, 2)
    record.update({
        "actual_base_cost": round(base_cost, 2),
        "actual_extra_cost": extra_cost if counted else 0.0,
        "actual_revenue": round(customer_revenue, 2) if counted else 0.0,
        "actual_profit": round(customer_revenue - total_cost, 2) if counted else 0.0,
        "cost_source": cost_source,
        "included_in_profit": realized,
    })
    record["raw_payload"] = _shipment_raw_payload(row, record)
    return record


def _shipment_raw_payload(row: list[Any], record: dict[str, Any]) -> list[Any]:
    # Keep the first 14 scraped columns + the included flag; the trailing slots
    # (legacy sheet-profit positions) are retained as None so older positional
    # readers don't shift. Real per-shipment economics live in the typed
    # actual_base_cost / actual_revenue / actual_profit columns.
    base = list(row[:14]) + [None] * max(0, 14 - len(row))
    included = "نعم" if record.get("included_in_profit", True) else "لا"
    return base[:14] + [included] + [None] * 9


def wallet_record(row: list[list[Any]], page: str, kind: str) -> dict[str, Any]:
    # Canonical wallet columns (WALLET_SPEC): 0=id 1=date 2=customer 3=type
    # 4=amount 5=description 6=tracking. Key on the unique transaction id (col 0)
    # so the row is stable even if other columns shift.
    return {
        "transaction_key": data_key(row, (0,)),
        "transaction_date": parse_date_text(row[1]) if len(row) > 1 else None,
        "user_name": norm(row[2]) if len(row) > 2 else "",
        "description": norm(row[5]) if len(row) > 5 else "",
        "amount": money(row[4]) if len(row) > 4 else 0.0,
        "transaction_type": norm(row[3]) if len(row) > 3 else kind,
        "balance_before": None,
        "balance_after": None,
        "source_page": page,
        "raw_payload": row,
    }


def payment_record(row: list[list[Any]]) -> dict[str, Any]:
    return {
        "payment_key": data_key(row, (0,)),
        "payment_date": parse_date_text(row[1]) if len(row) > 1 else None,
        "customer_name": norm(row[2]) if len(row) > 2 else "",
        "amount": money(row[4]) if len(row) > 4 else 0.0,
        "method": norm(row[3]) if len(row) > 3 else "",
        "status": "",
        "source_page": "wallet.php",
        "raw_payload": row,
    }


def pending_recharge_record(row: list[list[Any]]) -> dict[str, Any]:
    return {
        "recharge_key": data_key(row, (0,)),
        "customer_name": norm(row[1]) if len(row) > 1 else "",
        "amount": money(row[2]) if len(row) > 2 else 0.0,
        "reference": norm(row[3]) if len(row) > 3 else "",
        "status": norm(row[4]) if len(row) > 4 else "",
        "processed_by": norm(row[5]) if len(row) > 5 else "",
        "request_date": parse_date_text(row[6]) if len(row) > 6 else None,
        "raw_payload": row,
    }


def cod_record(row: list[list[Any]]) -> dict[str, Any]:
    # collect-cod.php has 5 columns: الشحنة | العميل | المبلغ | تاريخ التحصيل | تاريخ التحويل
    # (the order id is prefixed with '#', which we strip so it joins to shipments.order_id).
    return {
        "order_id": norm(row[0]).lstrip("#").strip() if len(row) > 0 else "",
        "tracking_number": "",
        "merchant_name": norm(row[1]) if len(row) > 1 else "",
        "customer_name": norm(row[1]) if len(row) > 1 else "",
        "cod_amount": money(row[2]) if len(row) > 2 else 0.0,
        "collection_date": parse_date_text(row[3]) if len(row) > 3 else None,
        "transfer_date": parse_date_text(row[4]) if len(row) > 4 else None,
        "settlement_status": "",
        "shipment_status": "",
        "raw_payload": row,
    }


# ---------------------------------------------------------------------------
# Header-based column mapping. We match each field by its Arabic header text and
# remap every table into a fixed canonical column order, so the scrape is
# resilient to the site reordering / renaming columns (positional parsing would
# silently corrupt data instead — that was the root cause of the COD bug).
# Each spec entry is (canonical_field, [header aliases]).
# ---------------------------------------------------------------------------

SHIPMENT_SPEC = [
    ("order_id", ["#", "الرقم", "رقم الطلب"]),
    ("tracking_number", ["رقم التتبع", "التتبع"]),
    ("merchant", ["العميل", "التاجر"]),
    ("store", ["المرسل", "المتجر"]),
    ("customer", ["المستلم", "العميل النهائي"]),
    ("city", ["المدينة"]),
    ("shipping_charge", ["التكلفة", "تكلفة الشحن", "قيمة الشحن"]),
    ("cod_amount", ["مبلغ COD", "مبلغ cod"]),
    ("order_amount", ["قيمة المنتجات", "قيمة الطلب"]),
    ("payment_type", ["نوع الدفع", "الدفع"]),
    ("carrier", ["شركة الشحن", "الناقل"]),
    ("weight", ["الوزن"]),
    ("status", ["الحالة"]),
    ("shipment_date", ["التاريخ", "تاريخ الشحن"]),
    ("delivery_date", ["تاريخ التسليم", "تاريخ التوصيل"]),
]

WALLET_SPEC = [
    ("id", ["الرقم"]),
    ("date", ["التاريخ"]),
    ("customer", ["العميل"]),
    ("type", ["النوع"]),
    ("amount", ["المبلغ"]),
    ("description", ["الوصف"]),
    ("tracking", ["رقم التتبع"]),
]

PENDING_RECHARGE_SPEC = [
    ("id", ["الرقم"]),
    ("customer", ["العميل"]),
    ("amount", ["المبلغ"]),
    ("reference", ["رقم المرجع", "المرجع"]),
    ("status", ["الحالة"]),
    ("processed_by", ["معالج بواسطة", "المعالج"]),
    ("date", ["التاريخ"]),
]

COD_SPEC = [
    ("order_id", ["الشحنة", "#", "رقم الطلب"]),
    ("customer", ["العميل"]),
    ("amount", ["المبلغ"]),
    ("collection_date", ["تاريخ التحصيل"]),
    ("transfer_date", ["تاريخ التحويل"]),
]


def _header_key(value: Any) -> str:
    text = re.sub(r"\s+", "", norm(value)).lower()
    return text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")


def remap_table(rows: list[list[Any]], spec: list[tuple[str, list[str]]]) -> list[list[Any]]:
    """Reorder a parsed table into the spec's canonical column order by matching
    each field's header aliases. Returns [canonical_header] + remapped data rows.
    Missing columns become None (not a wrong neighbouring column)."""
    if not rows:
        return rows
    header = {_header_key(h): i for i, h in enumerate(rows[0])}
    used: set[int] = set()
    col_for: list[int | None] = []
    for _field, aliases in spec:
        idx = None
        for alias in aliases:
            cand = header.get(_header_key(alias))
            if cand is not None and cand not in used:
                idx = cand
                break
        if idx is not None:
            used.add(idx)
        col_for.append(idx)
    out: list[list[Any]] = [[field for field, _ in spec]]
    for row in rows[1:]:
        out.append([(row[ci] if ci is not None and ci < len(row) else None) for ci in col_for])
    return out


def _largest_table(html: str) -> list[list[str]]:
    tables = parse_tables(html)
    return max(tables, key=len) if tables else []


def extract_shipments(html: str) -> tuple[list[list[Any]], int]:
    rows = remap_table([normalize_headers(r) for r in _largest_table(html)], SHIPMENT_SPEC)
    return rows, max(0, len(rows) - 1)


def extract_wallet(html: str) -> tuple[list[list[Any]], int, int]:
    tx_rows: list[list[Any]] = []
    for table in parse_tables(html):
        joined = " ".join(normalize_headers(table[0])) if table else ""
        if "الرقم" in joined and "الوصف" in joined:
            tx_rows = remap_table([normalize_headers(r) for r in table], WALLET_SPEC)
            break
    return tx_rows, max(0, len(tx_rows) - 1), 0


def extract_pending_recharges(html: str) -> list[list[Any]]:
    for table in parse_tables(html):
        joined = " ".join(normalize_headers(table[0])) if table else ""
        if "رقم المرجع" in joined and "معالج بواسطة" in joined:
            return remap_table([normalize_headers(r) for r in table], PENDING_RECHARGE_SPEC)
    return []


def extract_cod(html: str) -> list[list[Any]]:
    for table in parse_tables(html):
        joined = " ".join(normalize_headers(table[0])) if table else ""
        if "تاريخ التحصيل" in joined and "تاريخ التحويل" in joined:
            return remap_table([normalize_headers(r) for r in table], COD_SPEC)
    return []


def extract_shipping_companies(html: str, vat_rate: float = 0.15) -> list[dict[str, Any]]:
    """Parse /admin/shipping-companies.php into carrier price rows.

    Each carrier card exposes: a name (`🚚 ...`), the platform cost
    (`تكلفتك من المنصة: X ريال`), and an editable `customer_price` input. The
    site only shows gross prices, so we derive the ex-VAT net (gross / (1+vat))
    using the admin-configured VAT rate from the DB. Raw labels are mapped to the
    canonical carrier names via CARRIER_ALIASES (unknown carriers keep their
    raw label, so newly-added companies still flow through)."""
    if not html:
        return []
    from scripts.pricing import CARRIER_ALIASES

    def net(gross: float) -> float:
        return round(gross / (1 + vat_rate), 2)

    rows: list[dict[str, Any]] = []
    for m in re.finditer(r'<input\b[^>]*\bname="customer_price"[^>]*>', html):
        value_match = re.search(r'\bvalue="([0-9.]+)"', m.group(0))
        if not value_match:
            continue
        before = html[max(0, m.start() - 3200):m.start()]
        name_matches = re.findall(r'🚚\s*([^<>]{2,60}?)\s*<', before)
        plat_matches = re.findall(r'تكلفتك من المنصة[^0-9]{0,80}?([0-9]+(?:\.[0-9]+)?)\s*ريال', before)
        if not name_matches or not plat_matches:
            continue
        raw_name = norm(name_matches[-1])
        carrier_name = CARRIER_ALIASES.get(raw_name, raw_name)
        customer_gross = float(value_match.group(1))
        platform_gross = float(plat_matches[-1])
        if not customer_gross:
            continue
        rows.append({
            "carrier_name": carrier_name,
            "customer_gross": customer_gross,
            "customer_net": net(customer_gross),
            "platform_gross": platform_gross,
            "platform_net": net(platform_gross),
            "source": "scrape",
        })
    return rows


class _PlainText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        data = data.strip()
        if data:
            self.parts.append(data)


def _page_text(html: str) -> str:
    p = _PlainText()
    try:
        p.feed(html or "")
    except Exception:
        pass
    return re.sub(r"\s+", " ", " ".join(p.parts))


def _num_after(text: str, label: str):
    m = re.search(re.escape(label) + r"\s*([\d,]+(?:\.\d+)?)", text)
    return float(m.group(1).replace(",", "")) if m else None


def extract_site_kpis(html: dict[str, str]) -> dict[str, Any]:
    """The platform's own reported KPI cards (dashboard / reports / wallet), so the
    dashboard can show — and explain — where our aggregates differ from the site."""
    d = _page_text(html.get("dashboard.php", ""))
    r = _page_text(html.get("reports.php", ""))
    w = _page_text(html.get("wallet.php", ""))
    return {
        "total_orders": _num_after(d, "اجمالي الطلبات"),
        "delivered": _num_after(d, "تم التسليم"),
        "in_progress": _num_after(d, "قيد التنفيذ"),
        "customers": _num_after(d, "إجمالي العملاء"),
        "month_shipments": _num_after(r, "إجمالي الشحنات"),
        "month_revenue": _num_after(r, "إجمالي الإيرادات"),
        "month_realized_revenue": _num_after(r, "الإيرادات المحققة"),
        "month_base_cost": _num_after(r, "التكلفة الأساسية"),
        "month_net_profit": _num_after(r, "صافي الأرباح"),
        "deposits": _num_after(w, "إجمالي الإيداعات"),
    }


def _label_key(s: Any) -> str:
    return re.sub(r"\s+", "", norm(s))


_TOTAL_ROW_TOKENS = ("الإجمالي", "المجموع", "الاجمالي", "الكلي")


def extract_reports(html: str) -> dict[str, Any]:
    """Parse Lead's التقارير الشاملة (reports.php) — the platform's OWN actual
    revenue/cost/profit for the queried period. This is Lead's source of truth,
    so we never need a price sheet: `base_cost`/`net_profit` are the real carrier
    costs Lead paid, by construction (net_profit = realized_revenue - base_cost).

    Returns {totals, carriers, merchants}. carrier/merchant cost & profit are
    Lead's actuals (special-merchant pricing already baked in)."""
    text = _page_text(html)
    totals = {
        "revenue": _num_after(text, "إجمالي الإيرادات"),
        "realized_revenue": _num_after(text, "الإيرادات المحققة"),
        "base_cost": _num_after(text, "التكلفة الأساسية"),
        "net_profit": _num_after(text, "صافي الأرباح"),
    }
    carriers: list[dict[str, Any]] = []
    merchants: list[dict[str, Any]] = []

    def is_total_row(name: str) -> bool:
        return any(tok in name for tok in _TOTAL_ROW_TOKENS)

    for table in parse_tables(html):
        if not table or len(table) < 2:
            continue
        header = [_label_key(h) for h in table[0]]
        hset = set(header)
        # Per-carrier breakdown: شركة الشحن + التكلفة + صافي الربح.
        if _label_key("شركة الشحن") in hset and _label_key("صافي الربح") in hset:
            for row in table[1:]:
                name = norm(row[0]) if row else ""
                if not name or is_total_row(name):
                    continue
                carriers.append({
                    "carrier": name,
                    "count": int(money(row[1])) if len(row) > 1 else 0,
                    "revenue": money(row[2]) if len(row) > 2 else 0.0,
                    "cost": money(row[3]) if len(row) > 3 else 0.0,
                    "profit": money(row[4]) if len(row) > 4 else 0.0,
                })
        # Per-merchant breakdown: العميل (col 0) + الربح, but NOT صافي الربح.
        elif header and header[0] == _label_key("العميل") and _label_key("الربح") in hset:
            for row in table[1:]:
                name = norm(row[0]) if row else ""
                if not name or is_total_row(name):
                    continue
                revenue = money(row[3]) if len(row) > 3 else 0.0
                profit = money(row[4]) if len(row) > 4 else 0.0
                merchants.append({
                    "merchant": name,
                    "store": norm(row[1]) if len(row) > 1 else "",
                    "count": int(money(row[2])) if len(row) > 2 else 0,
                    "revenue": revenue,
                    "profit": profit,
                    "cost": round(revenue - profit, 2),
                })
    return {"totals": totals, "carriers": carriers, "merchants": merchants}


def extract_cod_fee(html: str) -> float | None:
    """The fixed COD fee from cod-settings.php (e.g. 5 ريال per COD order)."""
    if not html:
        return None
    m = re.search(r'name="cod_fee"[^>]*\bvalue="([0-9.]+)"', html)
    if m:
        return float(m.group(1))
    m = re.search(r"الرسوم الحالية[:：]?\s*([0-9.]+)", _page_text(html))
    return float(m.group(1)) if m else None


def fetch_reports(opener, base: str, start_date: str, end_date: str) -> str:
    """Fetch reports.php for an explicit date window (period=custom). Lead returns
    that period's frozen actuals, so any range — past or present — is exact."""
    url = (f"{base.rstrip('/')}/admin/reports.php?period=custom"
           f"&start_date={start_date}&end_date={end_date}")
    return safe_http_get(opener, url)


def _report_ranges(today: dt.date, floor: dt.date | None = None) -> list[tuple[str, str, str]]:
    """(label, start, end) windows to capture Lead's actual reports for. Labels are
    SEMANTIC ('all' / 'month' / '30d' / 'YYYY-MM') so the dashboard maps its presets
    to them without exact-date arithmetic — robust to the day rolling over between
    a capture and a page load. Closed months are frozen, so they stay stable.

    `floor` (the admin-set min_sync_date) anchors the 'all' window so all-time
    matches the data we actually keep — no pre-floor noise."""
    all_start = floor.isoformat() if floor else "2024-01-01"
    out: list[tuple[str, str, str]] = [
        ("all", all_start, today.isoformat()),
        ("month", today.replace(day=1).isoformat(), today.isoformat()),
        ("30d", (today - dt.timedelta(days=29)).isoformat(), today.isoformat()),
    ]
    y, m = today.year, today.month
    for _ in range(6):                                                       # last 6 closed months
        m -= 1
        if m == 0:
            m, y = 12, y - 1
        first = dt.date(y, m, 1)
        nxt = dt.date(y + (1 if m == 12 else 0), 1 if m == 12 else m + 1, 1)
        out.append((f"{y:04d}-{m:02d}", first.isoformat(), (nxt - dt.timedelta(days=1)).isoformat()))
    return out


LIST_CAP = 100  # the Lead admin list pages hard-cap at 100 rows


def window_scrape(opener, base, path, extract_rows, key_fn, date_indexes):
    """Capture EVERY row from a list page the server caps at LIST_CAP, by slicing
    the date range (month -> week -> day) until each window is under the cap and
    deduping by key_fn. Mirrors the sister project's window-scrape technique.

    extract_rows(html) -> list of rows (row[0] is the header).
    key_fn(row) -> dedup key. date_indexes -> columns holding the row's date
    (used only to anchor the walk at the most recent data, guarding clock skew).
    """
    seen: dict[str, list[Any]] = {}
    header: list[Any] = []

    def grab(a: dt.date, b: dt.date) -> list[list[Any]]:
        url = f"{base}/admin/{path}?date_from={a.isoformat()}&date_to={b.isoformat()}"
        rows = extract_rows(safe_http_get(opener, url))
        if rows:
            if not header:
                header.extend(rows[0])
            return rows[1:]
        return []

    def keep(rows: list[list[Any]]) -> None:
        for row in rows:
            key = key_fn(row)
            if key and key.strip("| "):
                seen[key] = row

    def split_days(a: dt.date, b: dt.date) -> None:
        day = a
        while day <= b:
            keep(grab(day, day))
            day += dt.timedelta(days=1)

    def handle(a: dt.date, b: dt.date) -> int:
        rows = grab(a, b)
        if len(rows) < LIST_CAP:
            keep(rows)
            return len(rows)
        start = a
        while start <= b:
            end = min(start + dt.timedelta(days=6), b)
            week = grab(start, end)
            if len(week) < LIST_CAP:
                keep(week)
            else:
                split_days(start, end)
            start = end + dt.timedelta(days=1)
        return len(rows)

    today = dt.date.today()
    # Anchor the walk at the most recent data so a lagging server clock can't
    # make us start "before" the newest rows and miss them.
    newest = today
    for row in (extract_rows(safe_http_get(opener, f"{base}/admin/{path}")) or [])[1:]:
        for idx in date_indexes:
            parsed = parse_date_text(row[idx]) if idx < len(row) else None
            if parsed and parsed > newest:
                newest = parsed

    year, month, empty = newest.year, newest.month, 0
    while True:
        first = dt.date(year, month, 1)
        last = dt.date(year + (month == 12), (month % 12) + 1, 1) - dt.timedelta(days=1)
        count = handle(first, last)
        empty = empty + 1 if count == 0 else 0
        if empty >= 4 or year < 2023:
            break
        month -= 1
        if month == 0:
            month, year = 12, year - 1

    return [header or []] + list(seen.values())


def normalize_row(values: list[Any]) -> list[Any]:
    return [None if v == "" else v for v in values]


def main() -> int:
    env = load_env(ENV_PATH)
    database_url_value = os.environ.get("DATABASE_URL", "")
    database_url_present = bool(database_url_value.strip())
    database_url_length = len(database_url_value.strip())
    db_module, db_import_error = load_db_store()
    # load_db_store() puts PROJECT_DIR on sys.path, so scripts.* is importable now.
    from scripts import pricing

    # Sync-floor date + VAT rate come from the DB (admin-editable), not hardcoded.
    min_sync_date = None
    carrier_vat_rate = 0.15
    if db_module is not None:
        try:
            with db_module.get_conn() as conn:
                raw_floor = db_module.get_setting(conn, "min_sync_date", "")
                if raw_floor.strip():
                    min_sync_date = dt.date.fromisoformat(raw_floor.strip())
                snap_settings = db_module.load_pricing_snapshot(conn).get("settings", {})
                if snap_settings.get("vat_rate") is not None:
                    carrier_vat_rate = float(snap_settings["vat_rate"])
        except Exception as exc:
            print(f"[sync] settings read skipped: {exc}", file=sys.stderr)

    for key in ("LEAD_USERNAME", "LEAD_PASSWORD", "LEAD_BASE_URL"):
        if not env.get(key):
            print(f"Missing {key} in .env", file=sys.stderr)
            return 2
    payload = scrape_site(env)
    source_type = "website"

    # Pure scrape → parse → upsert. No Excel staging: rows are parsed straight
    # from the scraped HTML, date-filtered, and handed to the DB upsert (which
    # dedupes on the natural keys).
    # Complete scrape: shipments.php and wallet.php are server-capped at 100 rows,
    # so we date-window-slice them to capture EVERY row (not just the first page).
    # collect-cod.php ignores date filters (a fixed current-state view), so it
    # stays a single fetch.
    base_url = env["LEAD_BASE_URL"].rstrip("/")
    opener = payload.get("opener")
    scrape_complete = opener is not None  # windowed (complete) scrape, not single-page
    if opener is not None:
        ship_rows = window_scrape(opener, base_url, "shipments.php",
                                  lambda h: extract_shipments(h)[0],
                                  lambda r: norm(r[0]) if r else "", (13,))
        wallet_rows = window_scrape(opener, base_url, "wallet.php",
                                    lambda h: extract_wallet(h)[0],
                                    lambda r: data_key(r, (0,)), (1,))
        pending_recharge_rows = window_scrape(opener, base_url, "pending-recharges.php",
                                              extract_pending_recharges,
                                              lambda r: data_key(r, (0,)), (6,))
    else:  # fallback: single-page fetch (e.g. opener unavailable)
        ship_rows = extract_shipments(payload["html"].get("shipments.php", ""))[0]
        wallet_rows = extract_wallet(payload["html"].get("wallet.php", ""))[0]
        pending_recharge_rows = extract_pending_recharges(payload["html"].get("pending-recharges.php", ""))
    wallet_total = max(0, len(wallet_rows) - 1)
    payments_total = wallet_total
    pending_recharge_total = max(0, len(pending_recharge_rows) - 1)
    cod_rows = extract_cod(payload["html"].get("collect-cod.php", ""))

    shipments_added = shipments_updated = shipments_dups = 0
    shipments_pruned = 0
    wallet_added = wallet_updated = wallet_dups = 0
    payments_added = payments_updated = payments_dups = 0
    pending_recharge_added = pending_recharge_updated = pending_recharge_dups = 0
    cod_added = cod_updated = cod_dups = 0

    if ship_rows:
        ship_rows = [ship_rows[0]] + [row for row in ship_rows[1:] if row_date_at_or_after(row, (13,), min_sync_date)]
        shipments_added = len(ship_rows) - 1
    if wallet_rows:
        wallet_rows = [wallet_rows[0]] + [row for row in wallet_rows[1:] if row_date_at_or_after(row, (1,), min_sync_date)]
        wallet_added = payments_added = len(wallet_rows) - 1
    if pending_recharge_rows:
        pending_recharge_rows = [pending_recharge_rows[0]] + [
            row for row in pending_recharge_rows[1:] if row_date_at_or_after(row, (6,), min_sync_date)
        ]
        pending_recharge_added = len(pending_recharge_rows) - 1
    if cod_rows:
        # collect-cod.php lists only a small, current set of COD orders, so capture
        # all of them (no date filter) — this is what carries the real collection /
        # transfer dates for every COD order, including older ones.
        cod_added = len(cod_rows) - 1

    postgres_enabled = bool(
        os.environ.get("DATABASE_URL", "").strip()
        or os.environ.get("POSTGRES_URL", "").strip()
        or os.environ.get("RAILWAY_DATABASE_URL", "").strip()
    )
    postgres_connected = False
    sync_run_id = None
    db_report = {"enabled": postgres_enabled, "connected": False, "synced": False, "comparison": {}}
    if postgres_enabled:
        if db_module is None:
            db_report = {
                "enabled": True,
                "connected": False,
                "synced": False,
                "error": f"{type(db_import_error).__name__}: {db_import_error}",
                "sync_run_id": None,
            }
        else:
            try:
                with db_module.get_conn() as conn:
                    postgres_connected = True
                    db_module.ensure_schema(conn)
                    sync_run_id = db_module.make_sync_run(conn, "sync_from_lead.py")
                    # Refresh carrier prices from the Lead site (non-fatal: a scrape
                    # hiccup must not zero out profits — fall back to seeded carriers).
                    try:
                        carrier_rows = extract_shipping_companies(payload["html"].get("shipping-companies.php", ""), carrier_vat_rate)
                        if carrier_rows:
                            db_module.upsert_carriers(conn, carrier_rows)
                    except Exception as carrier_exc:
                        print(f"[sync] carrier price refresh skipped: {carrier_exc}", file=sys.stderr)
                    snapshot = pricing.PricingSnapshot.from_db(db_module.load_pricing_snapshot(conn))
                    # Capture the platform's own KPI cards so the dashboard can show /
                    # explain where our aggregates differ from the site (non-fatal).
                    try:
                        site_kpis = extract_site_kpis(payload["html"])
                        site_kpis["captured_at"] = dt.datetime.now().isoformat(timespec="seconds")
                        db_module.set_setting(conn, "site_kpis", json.dumps(site_kpis, ensure_ascii=False))
                    except Exception as site_exc:
                        print(f"[sync] site KPI capture skipped: {site_exc}", file=sys.stderr)
                    # Capture Lead's OWN actual reports (revenue/cost/profit +
                    # per-carrier/-merchant) for the dashboard's common ranges, so
                    # the dashboard shows Lead's books, not a price-sheet estimate.
                    try:
                        if opener is not None:
                            # Start from the previous good snapshot so a failed fetch
                            # (e.g. the single-session Lead account got kicked by a
                            # concurrent browser login) can't blank out the actuals —
                            # we only ever REPLACE a range we successfully re-fetched.
                            prev = {}
                            try:
                                raw_prev = db_module.get_setting(conn, "lead_reports", "")
                                if raw_prev:
                                    prev = json.loads(raw_prev)
                            except Exception:
                                prev = {}
                            snaps: dict[str, Any] = dict(prev.get("ranges", {}))
                            captured = 0
                            for r_label, r_start, r_end in _report_ranges(dt.date.today(), min_sync_date):
                                parsed = extract_reports(fetch_reports(opener, base_url, r_start, r_end))
                                if parsed["totals"].get("revenue") is not None:
                                    parsed["range"] = {"start": r_start, "end": r_end}
                                    snaps[r_label] = parsed
                                    captured += 1
                            cod_html = safe_http_get(opener, f"{base_url}/admin/cod-settings.php")
                            cod_fee = extract_cod_fee(cod_html) or prev.get("cod_fee")
                            # Only persist if we actually captured something this run
                            # (else keep the prior good snapshot untouched).
                            if captured:
                                db_module.set_setting(conn, "lead_reports", json.dumps({
                                    "ranges": snaps, "cod_fee": cod_fee,
                                    "captured_at": dt.datetime.now().isoformat(timespec="seconds"),
                                }, ensure_ascii=False))
                            else:
                                print("[sync] lead reports: 0 ranges captured — kept previous snapshot", file=sys.stderr)
                    except Exception as rep_exc:
                        print(f"[sync] lead reports capture skipped: {rep_exc}", file=sys.stderr)
                    # Exact per-shipment cost from the billed invoices (non-fatal:
                    # if download/parse fails, the unbilled path computes from the
                    # live carrier prices for everyone).
                    invoice_costs: dict[str, Any] = {}
                    try:
                        if opener is not None:
                            invoice_costs = collect_invoice_costs(opener, base_url)
                    except Exception as inv_exc:
                        print(f"[sync] invoice cost collection skipped: {inv_exc}", file=sys.stderr)
                    invoiced_n = sum(1 for r in ship_rows[1:] if norm(r[0]) in invoice_costs)
                    print(f"[sync] invoice costs: {len(invoice_costs)} billed shipments, {invoiced_n} match this scrape", file=sys.stderr)
                    detected_tax_agreements = wallet_tax_agreement_customers(wallet_rows[1:])
                    detected_tax_agreements.update(stored_wallet_tax_agreement_customers(conn))
                    detected_tax_agreements.update(bank_vat_agreement_customers())
                    if detected_tax_agreements:
                        db_module.upsert_customer_tax_agreements(conn, sorted(detected_tax_agreements), "wallet_or_bank_vat_deduction")
                    tax_agreement_customers = db_module.load_customer_tax_agreements(conn)
                    print(f"[sync] customer tax agreements: {len(tax_agreement_customers)}", file=sys.stderr)
                    shipments_payload = [
                        shipment_record(row, snapshot, invoice_costs, tax_agreement_customers, carrier_vat_rate)
                        for row in ship_rows[1:]
                    ]
                    wallet_payload = [wallet_record(row, "wallet.php", "transaction") for row in wallet_rows[1:]]
                    payments_payload = [payment_record(row) for row in wallet_rows[1:]]
                    pending_recharge_payload = [pending_recharge_record(row) for row in pending_recharge_rows[1:]]
                    cod_payload = [cod_record(row) for row in cod_rows[1:]]
                    shipment_ins, shipment_upd = db_module.upsert_rows(conn, "shipments", shipments_payload, ["order_id"], [
                        "tracking_number", "merchant_name", "store_name", "customer_name", "city", "carrier", "payment_type",
                        "status", "shipment_date", "delivery_date", "weight", "cod_amount", "shipping_charge",
                        "included_in_profit", "actual_base_cost", "actual_extra_cost", "actual_revenue", "actual_profit", "cost_source",
                        "source_hash", "raw_payload",
                    ])
                    # Prune shipments the site no longer returns (deleted/draft orders),
                    # but ONLY after a healthy complete scrape: require the windowed
                    # path AND that we captured at least 90% of the rows already stored,
                    # so a partial/failed scrape can never delete valid data.
                    shipments_pruned = 0
                    kept_ids = [norm(r[0]) for r in ship_rows[1:] if r and norm(r[0])]
                    if scrape_complete and len(kept_ids) >= 100:
                        existing = db_module.compare_counts(conn)["shipments"]
                        if existing == 0 or len(kept_ids) >= 0.9 * existing:
                            shipments_pruned = db_module.prune_missing_shipments(conn, kept_ids)
                    wallet_ins, wallet_upd = db_module.upsert_rows(conn, "wallet_transactions", wallet_payload, ["transaction_key"], [
                        "transaction_date", "user_name", "description", "amount", "transaction_type", "balance_before",
                        "balance_after", "source_page", "raw_payload",
                    ])
                    payment_ins, payment_upd = db_module.upsert_rows(conn, "payments", payments_payload, ["payment_key"], [
                        "payment_date", "customer_name", "amount", "method", "status", "source_page", "raw_payload",
                    ])
                    pending_recharge_ins, pending_recharge_upd = db_module.upsert_rows(conn, "pending_recharges", pending_recharge_payload, ["recharge_key"], [
                        "request_date", "customer_name", "amount", "reference", "status", "processed_by", "raw_payload",
                    ])
                    pending_recharge_updated = pending_recharge_upd
                    cod_ins, cod_upd = db_module.upsert_rows(conn, "cod_collections", cod_payload, ["order_id"], [
                        "tracking_number", "merchant_name", "customer_name", "cod_amount", "collection_date",
                        "transfer_date", "settlement_status", "shipment_status", "raw_payload",
                    ])
                    comparison = db_module.compare_counts(conn)
                    db_module.finish_sync_run(
                        conn,
                        sync_run_id,
                        "ok",
                        inserted=shipment_ins + wallet_ins + payment_ins + pending_recharge_ins + cod_ins,
                        updated=shipment_upd + wallet_upd + payment_upd + pending_recharge_upd + cod_upd,
                        skipped=shipments_dups + wallet_dups + payments_dups + pending_recharge_dups + cod_dups,
                    )
                    db_report = {
                        "enabled": True,
                        "connected": True,
                        "synced": True,
                        "comparison": comparison,
                        "db_rows_inserted": shipment_ins + wallet_ins + payment_ins + pending_recharge_ins + cod_ins,
                        "db_rows_updated": shipment_upd + wallet_upd + payment_upd + pending_recharge_upd + cod_upd,
                        "db_rows_skipped": shipments_dups + wallet_dups + payments_dups + pending_recharge_dups + cod_dups,
                        "sync_run_id": sync_run_id,
                    }
            except Exception as exc:
                db_report = {
                    "enabled": True,
                    "connected": postgres_connected,
                    "synced": False,
                    "error": str(exc),
                    "sync_run_id": sync_run_id,
                }

    # Profit is now computed inline (above) from the DB pricing snapshot, and the
    # dashboard is rendered live from PostgreSQL by app.py — so the legacy
    # Excel report build and the static site-generation subprocesses are gone.
    build_report_error = None
    site_error = None
    build_completed = None
    site_completed = None

    report = {
        "source_type": source_type,
        "shipments_rows": max(0, len(ship_rows) - 1),
        "shipments_new": shipments_added,
        "shipments_updated": shipments_updated,
        "shipments_duplicates": shipments_dups,
        "wallet_rows": wallet_total,
        "wallet_new": wallet_added,
        "wallet_updated": wallet_updated,
        "wallet_duplicates": wallet_dups,
        "payments_rows": payments_total,
        "payments_new": payments_added,
        "payments_updated": payments_updated,
        "payments_duplicates": payments_dups,
        "pending_recharge_rows": pending_recharge_total,
        "pending_recharge_new": pending_recharge_added,
        "pending_recharge_updated": pending_recharge_updated,
        "pending_recharge_duplicates": pending_recharge_dups,
        "cod_rows": max(0, len(cod_rows) - 1),
        "cod_new": cod_added,
        "cod_updated": cod_updated,
        "cod_duplicates": cod_dups,
        "errors": 0,
        "api_observations": len(payload.get("logs", [])),
        "checks": payload.get("checks", {}),
        "login_status": payload.get("checks", {}).get("login_status"),
        "login_final_url": payload.get("checks", {}).get("login_final_url"),
        "login_http_status": payload.get("checks", {}).get("login_http_status"),
        "login_redirected_to_login": payload.get("checks", {}).get("login_redirected_to_login"),
        "login_html_500": payload.get("checks", {}).get("login_html_500"),
        "postgres": db_report,
        "postgres_enabled": postgres_enabled,
        "postgres_connected": postgres_connected,
        "rows_inserted": db_report.get("db_rows_inserted", 0),
        "rows_updated": db_report.get("db_rows_updated", 0),
        "rows_skipped": db_report.get("db_rows_skipped", 0),
        "sync_run_id": sync_run_id,
        "dashboard_refreshed": True,
        "build_report_error": build_report_error,
        "site_error": site_error,
        "build_report_exit_code": build_completed.returncode if build_completed is not None else None,
        "site_exit_code": site_completed.returncode if site_completed is not None else None,
    }

    state = load_state()
    state["last_run_at"] = dt.datetime.now().isoformat(timespec="seconds")
    state["network_events"] = len(payload.get("logs", []))
    save_state(state)
    login_ok = bool(
        payload.get("checks", {}).get("has_dashboard_links")
        and payload.get("checks", {}).get("shipments_opened")
        and payload.get("checks", {}).get("wallet_opened")
        and payload.get("checks", {}).get("cod_opened")
    )
    final_url = payload.get("checks", {}).get("dashboard_final_url") or payload.get("checks", {}).get("login_final_url") or ""
    sync_summary = {
        "source_type": source_type,
        "login_ok": login_ok,
        "final_url": final_url,
        "postgres_enabled": postgres_enabled,
        "postgres_connected": postgres_connected,
        "sync_run_id": sync_run_id,
        "rows_inserted": db_report.get("db_rows_inserted", 0),
        "rows_updated": db_report.get("db_rows_updated", 0),
        "rows_skipped": db_report.get("db_rows_skipped", 0),
        "shipments_rows": max(0, len(ship_rows) - 1),
        "shipments_pruned": shipments_pruned,
        "wallet_rows": wallet_total,
        "payments_rows": payments_total,
        "cod_rows": max(0, len(cod_rows) - 1),
        "exit_code": 0,
    }
    print(f"DATABASE_URL_PRESENT={str(database_url_present).lower()} DATABASE_URL_LENGTH={database_url_length}", flush=True)
    print(f"SYNC_SUMMARY {json.dumps(sync_summary, ensure_ascii=False)}", flush=True)
    print(f"[lead-sync] source={source_type}", flush=True)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
